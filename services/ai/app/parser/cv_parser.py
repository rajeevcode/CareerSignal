import io
import json
import re
from typing import Any

import fitz
from pydantic import BaseModel, Field

from app.config import settings
from app.matching.engine import generate_embedding
from app.matching.skill_graph import extract_skills_from_text, normalize_skills
from app.budget import BudgetExceeded
from app.openai_client import chat_completion


class ParsedExperience(BaseModel):
    company: str
    title: str
    start_date: str | None = None
    end_date: str | None = None
    location: str | None = None
    description: str | None = None
    achievements: list[str] = Field(default_factory=list)
    skills_used: list[str] = Field(default_factory=list)


class ParsedEducation(BaseModel):
    institution: str
    degree: str | None = None
    field: str | None = None
    end_date: str | None = None


class ParsedCV(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    location: str | None = None
    current_title: str | None = None
    total_experience_years: float | None = None
    profile_summary: str | None = None
    experience: list[ParsedExperience] = Field(default_factory=list)
    education: list[ParsedEducation] = Field(default_factory=list)
    skills: list[str] = Field(default_factory=list)
    achievements: list[str] = Field(default_factory=list)
    seniority: str | None = None
    gaps: list[str] = Field(default_factory=list)
    weak_sections: list[str] = Field(default_factory=list)


ALLOWED_CV_EXTENSIONS = (".pdf", ".docx")


class UnsupportedFileType(ValueError):
    """Raised when an uploaded CV is not a PDF or DOCX file."""


def extract_text_from_pdf(content: bytes) -> str:
    doc = fitz.open(stream=content, filetype="pdf")
    text_parts = []
    for page in doc:
        text_parts.append(page.get_text())
    doc.close()
    return "\n".join(text_parts).strip()


def extract_text_from_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def compute_ats_score(text: str, parsed: ParsedCV) -> float:
    score = 50.0
    if parsed.name:
        score += 5
    if parsed.email:
        score += 5
    if parsed.current_title:
        score += 5
    if parsed.experience:
        score += min(15, len(parsed.experience) * 5)
    if parsed.skills:
        score += min(10, len(parsed.skills))
    if parsed.profile_summary and len(parsed.profile_summary) > 50:
        score += 10
    bullet_count = len(re.findall(r"^[\•\-\*]", text, re.MULTILINE))
    if bullet_count >= 5:
        score += 10
    sections = ["experience", "education", "skills", "summary"]
    text_lower = text.lower()
    for section in sections:
        if section in text_lower:
            score += 2
    return min(100.0, score)


async def parse_cv_with_llm(text: str) -> ParsedCV:
    if not settings.openai_api_key:
        return _heuristic_parse(text)

    prompt = f"""Extract structured CV data from this resume text. Return JSON matching this schema:
{{
  "name", "email", "phone", "location", "current_title", "total_experience_years",
  "profile_summary", "experience": [{{company, title, start_date, end_date, location, description, achievements[], skills_used[]}}],
  "education": [{{institution, degree, field, end_date}}],
  "skills": [], "achievements": [], "seniority": "junior|mid|senior|lead",
  "gaps": [], "weak_sections": []
}}

Resume:
{text[:12000]}"""

    try:
        response = chat_completion(
            messages=[
                {"role": "system", "content": "Extract CV data accurately. Return valid JSON only."},
                {"role": "user", "content": prompt},
            ],
            response_format={"type": "json_object"},
            temperature=0.1,
            feature="parse_cv",
        )
        data = json.loads(response.choices[0].message.content or "{}")
        parsed = ParsedCV.model_validate(data)
        parsed.skills = normalize_skills(parsed.skills + extract_skills_from_text(text))
        return parsed
    except BudgetExceeded:
        raise
    except Exception:
        return _heuristic_parse(text)


def _heuristic_parse(text: str) -> ParsedCV:
    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone_match = re.search(r"\+?[\d\s\-()]{10,}", text)
    skills = extract_skills_from_text(text)
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    name = lines[0] if lines else None
    return ParsedCV(
        name=name,
        email=email_match.group() if email_match else None,
        phone=phone_match.group() if phone_match else None,
        skills=skills,
        profile_summary=text[:500] if text else None,
    )


async def parse_cv(content: bytes, filename: str) -> dict[str, Any]:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(content)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(content)
    else:
        raise UnsupportedFileType("Only PDF and DOCX files are supported.")

    parsed = await parse_cv_with_llm(text)
    ats_score = compute_ats_score(text, parsed)
    skills = normalize_skills(parsed.skills)

    summary_text = parsed.profile_summary or " ".join(
        [f"{e.title} at {e.company}" for e in parsed.experience[:3]]
    )
    embedding = await generate_embedding(f"{parsed.current_title or ''} {summary_text} {' '.join(skills)}")

    return {
        "parsed": parsed.model_dump(),
        "parsed_text": text,
        "ats_score": ats_score,
        "skills": skills,
        "embedding": embedding,
    }
